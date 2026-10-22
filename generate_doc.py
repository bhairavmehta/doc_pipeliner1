"""
Generate a Word + Markdown document from a content file (content.json).

Usage:
    pip install python-docx
    python generate_doc.py                      # uses content.json, writes Document.docx/.md
    python generate_doc.py content.json MyDoc   # custom content file and output name

Edit the wording in content.json; this script only handles layout.
Block types in content.json: heading1, heading2, paragraph (style "note" = grey italic),
bullets, table, code, image.
"""
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).parent
FONT = "Calibri"
NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = RGBColor(0x2F, 0x54, 0x96)
GREY_TXT = RGBColor(0x59, 0x59, 0x59)
FOOT_TXT = RGBColor(0x80, 0x80, 0x80)
HEADER_FILL = "1F3864"
BAND_FILL = "F2F2F2"
BORDER = "BFBFBF"


# ---------- low-level helpers ----------
def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_table_borders(table, color=BORDER, size=4):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:color"), color)
        borders.append(el)
    tbl_pr.append(borders)


def row_flag(row, tag):
    """tag = 'tblHeader' (repeat on each page) or 'cantSplit' (keep row on one page)."""
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement(f"w:{tag}")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def add_page_number(paragraph, size):
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    run.font.size = size
    run.font.color.rgb = FOOT_TXT


def style_run(run, size=10.5, bold=False, italic=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


# ---------- document setup ----------
def setup(doc, cfg, base):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin, sec.bottom_margin = Inches(0.9), Inches(0.75)
    sec.left_margin = sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    for name, size, color in (("Heading 1", 14, NAVY), ("Heading 2", 11.5, BLUE)):
        st = doc.styles[name]
        st.font.name, st.font.size, st.font.bold, st.font.color.rgb = FONT, Pt(size), True, color
        st.element.rPr.rFonts.set(qn("w:asciiTheme"), "") if False else None

    # optional logo in header (top-left, every page) - only if "logo" is set in the content file
    logo = cfg.get("logo")
    if logo and (base / logo).exists():
        sec.header.paragraphs[0].add_run().add_picture(str(base / logo), width=Inches(1.15))

    # footer text + page number (right-aligned)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_run(fp.add_run(cfg["footer"] + " · Page "), size=8, color=FOOT_TXT)
    add_page_number(fp, Pt(8))

    # title block
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    style_run(p.add_run(cfg["title"]), size=18, bold=True, color=NAVY)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    style_run(p.add_run(cfg["subtitle"]), size=11, color=GREY_TXT)
    p_pr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "8"), ("space", "4"), ("color", HEADER_FILL)):
        bottom.set(qn(f"w:{k}"), v)
    bdr.append(bottom)
    p_pr.append(bdr)


# ---------- block renderers ----------
def add_table(doc, b):
    cols, rows, widths = b["columns"], b["rows"], b["widths_in"]
    t = doc.add_table(rows=1 + len(rows), cols=len(cols))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    set_table_borders(t)
    layout = OxmlElement("w:tblLayout")          # fixed column widths (Word + LibreOffice)
    layout.set(qn("w:type"), "fixed")
    t._tbl.tblPr.append(layout)
    for i, col in enumerate(t.columns):
        col.width = Inches(widths[i])
    for r_i, values in enumerate([cols] + rows):
        row = t.rows[r_i]
        row_flag(row, "tblHeader" if r_i == 0 else "cantSplit")
        for c_i, text in enumerate(values):
            cell = row.cells[c_i]
            cell.width = Inches(widths[c_i])
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            run = para.add_run(text)
            if r_i == 0:
                style_run(run, size=9, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
                shade(cell, HEADER_FILL)
            else:
                style_run(run, size=9, bold=b.get("bold_first_column") and c_i == 0)
                if r_i % 2 == 0:
                    shade(cell, BAND_FILL)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def render_docx(cfg, base, out_path):
    doc = Document()
    setup(doc, cfg, base)
    for b in cfg["blocks"]:
        kind = b["type"]
        if kind in ("heading1", "heading2"):
            h = doc.add_heading(b["text"], level=1 if kind == "heading1" else 2)
            h.paragraph_format.keep_with_next = True
            h.paragraph_format.space_before = Pt(14 if kind == "heading1" else 10)
        elif kind == "paragraph":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            if b.get("style") == "note":
                style_run(p.add_run(b["text"]), size=9.5, italic=True, color=GREY_TXT)
            else:
                style_run(p.add_run(b["text"]))
        elif kind == "bullets":
            for item in b["items"]:
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_after = Pt(3)
                style_run(p.add_run(item))
        elif kind == "table":
            add_table(doc, b)
        elif kind == "code":
            for line in b["lines"]:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.left_indent = Inches(0.25)
                p_pr = p._p.get_or_add_pPr()
                shd = OxmlElement("w:shd")
                for k, v in (("val", "clear"), ("color", "auto"), ("fill", BAND_FILL)):
                    shd.set(qn(f"w:{k}"), v)
                p_pr.append(shd)
                r = p.add_run(line or " ")
                style_run(r, size=9.5)
                r.font.name = "Consolas"
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif kind == "image":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(base / b["file"]), width=Inches(b["width_in"]))
            c = doc.add_paragraph()
            c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            c.paragraph_format.space_after = Pt(8)
            style_run(c.add_run(b["caption"]), size=9, italic=True, color=GREY_TXT)
    doc.save(out_path)


def render_md(cfg, out_path):
    esc = lambda s: str(s).replace("|", "\\|")
    lines = [f"![Logo]({cfg['logo']})", ""] if cfg.get("logo") else []
    lines += [f"# {cfg['title']}", "", f"*{cfg['subtitle']}*", ""]
    for b in cfg["blocks"]:
        kind = b["type"]
        if kind == "heading1":
            lines += [f"## {b['text']}", ""]
        elif kind == "heading2":
            lines += [f"### {b['text']}", ""]
        elif kind == "paragraph":
            lines += [f"*{b['text']}*" if b.get("style") == "note" else b["text"], ""]
        elif kind == "bullets":
            lines += [f"- {i}" for i in b["items"]] + [""]
        elif kind == "table":
            lines.append("| " + " | ".join(b["columns"]) + " |")
            lines.append("|" + "|".join("---" for _ in b["columns"]) + "|")
            for r in b["rows"]:
                cells = [f"**{esc(c)}**" if b.get("bold_first_column") and i == 0 and c else (esc(c) or " ")
                         for i, c in enumerate(r)]
                lines.append("| " + " | ".join(cells) + " |")
            lines.append("")
        elif kind == "code":
            lines += ["```"] + b["lines"] + ["```", ""]
        elif kind == "image":
            lines += [f"![{b['caption']}]({b['file']})", "", f"*{b['caption']}*", ""]
    Path(out_path).write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    content = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "content.json"
    name = sys.argv[2] if len(sys.argv) > 2 else "Document"
    cfg = json.loads(content.read_text(encoding="utf-8"))
    base = content.parent
    render_docx(cfg, base, base / f"{name}.docx")
    render_md(cfg, base / f"{name}.md")
    print(f"Wrote {name}.docx and {name}.md in {base}")
