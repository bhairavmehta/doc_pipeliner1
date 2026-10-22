"""
run_pipeline.py – from a user-story prompt to the finished requirements document and a run report,
using Codex (codex exec) as the AI and the Python scripts in this folder for everything else.

YOU PROVIDE
  --story     the user story prompt (a .md or .txt file with the story text, exactly as given to you)
  --research  research_findings.md (the research summary)
  --sources   folder with the source papers (PDF / HTML / TXT)          [optional but recommended]
  --inputs    folder with other supporting files (e.g. earlier deliverables: .md, .txt, .docx, .pdf) [optional]

IT PRODUCES  (in stories/<story id>/)
  facts.jsonl / facts.md     fact ledger: the findings split into numbered facts (K1)        (Python)
  plan.json                  document plan: outline, section types, fact assignments, figures (Codex)
  section_briefs.md          per-section brief: checklist, floor, assigned facts, gold rows   (Python)
  research/research_evidence.md  passages from the papers, cited by file and page           (Python)
  figures.py + fig*.png      the figures                                                     (Codex + Python)
  content.md / content.json  the document text                                               (Codex + Python)
  scorecard.json, findings.jsonl, critique_report.md   critic panel + quality gate (D1–D10)  (Python [+ SLM])
  figure_qa.json             measured figure defects (overlaps, overflow, arrows over text)  (Python)
  <output_name>.docx / .md   the requirements document                                       (Python)
  qa_report.md               fact-check and coverage review                                  (Codex)
  run_report.md / .docx      what happened in this run: steps, files, tokens, open items     (Python)
  logs/, _prompts/, run_state.json   logs, the exact prompts sent to Codex, step status

USAGE
  python run_pipeline.py --story story.md --research research_findings.md --sources papers/
  python run_pipeline.py --story story.md --research research_findings.md --dry-run   # write prompts only
  python run_pipeline.py --id 123456 --from draft      # re-run from a step (story already set up)
  python run_pipeline.py --id 123456 --only qa         # run one step
  Revise after review: write reviewer comments in stories/<id>/feedback.md, then  --id <id> --from draft
  Steps: check, setup, facts, plan, evidence, figures, draft, critique, build, qa, report

QUALITY OPTIONS
  --gold FILE         gold document the draft is scored against (default: examples/gold/US311617/…_benchmark.md)
  --draft-mode MODE   whole (one Codex call, default) or sections (one call per section, then assembled)
  --revise-rounds N   critique → revise rounds when the gate fails (default 2)
  --strict-gates      stop the run if the hard gates still fail after the revise rounds
  --critic-mode MODE  rules (default) or hybrid (rules + small model; needs SLM_* settings, see quality/slm.py)
  --judge             add the J1 pairwise judge against the gold document (needs JUDGE_* or LLM_* settings)
  --no-qa-fix         keep QA advisory (default: High QA issues get one fix round before the report)
  --validate-only     with --only plan / --only critique: check and score without calling Codex

CODEX SETTINGS
  Codex CLI must be installed and signed in (run `codex` once by hand first).
  CODEX_BIN    path to codex if it is not on PATH
  CODEX_ARGS   extra arguments for every codex call, e.g.  -m <model>
  Codex runs with --sandbox workspace-write in this folder (read-only for the QA step).
  QA_REVIEWER  who the QA step reviews as (or --reviewer); default: a subject-matter expert for the story's topic
"""
import argparse
import datetime as dt
import json
import os
import re
import shlex
import shutil
import subprocess
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent
PROMPTS = ROOT / "prompts"
STEPS = ["check", "setup", "facts", "plan", "evidence", "figures", "draft", "critique", "build", "qa", "report"]
CODEX_TIMEOUT = int(os.environ.get("CODEX_TIMEOUT", "1800"))  # seconds per Codex call
MAX_FIX_ROUNDS = 2
DEFAULT_GOLD = ROOT / "examples" / "gold" / "US311617" / "US311617_benchmark.md"

sys.path.insert(0, str(ROOT))


# ------------------------------------------------------------------ helpers
class StepError(Exception):
    pass


def now():
    return dt.datetime.now().isoformat(timespec="seconds")


def rel(p):
    try:
        return str(Path(p).resolve().relative_to(ROOT)).replace("\\", "/")
    except ValueError:
        return str(p)


class Run:
    def __init__(self, story_dir, dry_run=False, codex_args=None):
        self.dir = story_dir
        self.dry = dry_run
        self.codex_args = codex_args or []
        self.state_file = story_dir / "run_state.json"
        self.state = json.loads(self.state_file.read_text(encoding="utf-8")) if self.state_file.exists() else {"steps": {}}
        (story_dir / "logs").mkdir(parents=True, exist_ok=True)
        (story_dir / "_prompts").mkdir(parents=True, exist_ok=True)

    def save(self):
        self.state_file.write_text(json.dumps(self.state, indent=2), encoding="utf-8")

    def record(self, step, status, seconds, files=(), tokens=None, note=""):
        self.state["steps"][step] = {"status": status, "seconds": round(seconds, 1), "finished": now(),
                                     "files": [rel(f) for f in files], "tokens": tokens or {}, "note": note}
        self.save()

    # -------------------------------------------------------------- prompts
    def prompt(self, template, **extra):
        text = (PROMPTS / template).read_text(encoding="utf-8")
        values = {"STORY_DIR": rel(self.dir), "TODAY": dt.date.today().strftime("%B %d, %Y").replace(" 0", " ")}
        values.update(extra)
        for k, v in values.items():
            text = text.replace("{" + k + "}", str(v))
        return text

    # -------------------------------------------------------------- codex
    def codex(self, step, prompt_text, sandbox="workspace-write", out_name=None):
        """Run `codex exec` in the project folder with the prompt on stdin. Returns (final_message, tokens)."""
        (self.dir / "_prompts" / f"{step}.md").write_text(prompt_text, encoding="utf-8")
        if self.dry:
            return "", {}
        exe = os.environ.get("CODEX_BIN") or shutil.which("codex")
        if not exe:
            raise StepError("Codex CLI not found. Install it, sign in once with `codex`, or set CODEX_BIN.")
        last = self.dir / "logs" / (out_name or f"{step}_last_message.md")
        cmd = [exe, "exec", "--sandbox", sandbox, "--skip-git-repo-check", "--json",
               "--output-last-message", str(last), *self.codex_args, "-"]
        log = self.dir / "logs" / f"{step}.log"
        print(f"    codex exec ({sandbox}) ... log: {rel(log)}")
        try:
            proc = subprocess.run(cmd, cwd=ROOT, input=prompt_text.encode("utf-8"), capture_output=True,
                                  timeout=CODEX_TIMEOUT)
        except subprocess.TimeoutExpired:
            raise StepError(f"Codex did not finish within {CODEX_TIMEOUT} s (set CODEX_TIMEOUT to allow longer).")
        stdout = proc.stdout.decode("utf-8", "replace")
        stderr = proc.stderr.decode("utf-8", "replace")
        with log.open("a", encoding="utf-8") as f:
            f.write(f"\n===== {now()}  {' '.join(shlex.quote(c) for c in cmd)}\n--- stdout (JSON events)\n{stdout}\n--- stderr\n{stderr}\n")
        if proc.returncode != 0:
            tail = (stderr or stdout).strip().splitlines()[-5:]
            raise StepError(f"codex exec failed (exit {proc.returncode}): " + " | ".join(tail))
        return (last.read_text(encoding="utf-8") if last.exists() else ""), token_usage(stdout)


def token_usage(jsonl):
    """Sum token counts from codex --json events (any event carrying a 'usage' object)."""
    total = {}
    for line in jsonl.splitlines():
        line = line.strip()
        if not line.startswith("{"):
            continue
        try:
            ev = json.loads(line)
        except json.JSONDecodeError:
            continue
        stack = [ev]
        while stack:
            o = stack.pop()
            if isinstance(o, dict):
                u = o.get("usage")
                if isinstance(u, dict):
                    for k, v in u.items():
                        if isinstance(v, (int, float)):
                            total[k] = total.get(k, 0) + v
                stack.extend(v for k, v in o.items() if k != "usage")
            elif isinstance(o, list):
                stack.extend(o)
    return total


def add_tokens(a, b):
    for k, v in (b or {}).items():
        a[k] = a.get(k, 0) + v
    return a


def to_text(src, dst_dir):
    """Convert a supporting file to plain text for Codex (md/txt copied, pdf/docx/html extracted)."""
    dst_dir.mkdir(parents=True, exist_ok=True)
    out = dst_dir / (src.name + ".txt")
    suf = src.suffix.lower()
    try:
        if suf in (".md", ".txt", ".csv", ".json"):
            text = src.read_text(encoding="utf-8", errors="ignore")
        elif suf == ".pdf":
            from pypdf import PdfReader
            text = "\n\n".join(f"[page {i + 1}]\n{p.extract_text() or ''}" for i, p in enumerate(PdfReader(str(src)).pages))
        elif suf == ".docx":
            import docx
            d = docx.Document(str(src))
            parts = [p.text for p in d.paragraphs]
            for t in d.tables:
                for r in t.rows:
                    parts.append(" | ".join(c.text for c in r.cells))
            text = "\n".join(parts)
        elif suf in (".html", ".htm"):
            from research_findings import _HTMLText
            h = _HTMLText()
            h.feed(src.read_text(encoding="utf-8", errors="ignore"))
            text = "".join(h.parts)
        else:
            return None
    except Exception as e:
        print(f"    [skip] {src.name}: {e}")
        return None
    out.write_text(text, encoding="utf-8")
    return out


def copy_tree(src, dst):
    dst.mkdir(parents=True, exist_ok=True)
    n = 0
    for f in Path(src).rglob("*"):
        if f.is_file() and "_text" not in f.relative_to(src).parts:
            target = dst / f.relative_to(src)
            target.parent.mkdir(parents=True, exist_ok=True)
            if not target.exists() or target.stat().st_size != f.stat().st_size:
                shutil.copy2(f, target)
                n += 1
    return n


# ------------------------------------------------------------------ checks on content.md
def check_content(story_dir, plan):
    """Return a list of problems in content.md (empty list = OK). Also writes content.json when OK."""
    import md_to_content
    md = story_dir / "content.md"
    if not md.exists():
        return ["content.md was not created."]
    text = md.read_text(encoding="utf-8")
    problems = []
    if not text.lstrip().startswith("---"):
        problems.append("The file must start with the front matter block (--- / title: / subtitle: / footer: / ---).")
    if text.lstrip().startswith("```"):
        problems.append("Remove the ``` code fence around the whole document; the file must be plain Markdown.")
    cfg = md_to_content.parse(text)
    for key in ("title", "subtitle", "footer"):
        if not cfg.get(key) or "<" in cfg.get(key, ""):
            problems.append(f"Front matter '{key}' is missing or still a placeholder.")
    heads = [b["text"] for b in cfg["blocks"] if b["type"] == "heading1"]
    if len(heads) < 5:
        problems.append(f"Only {len(heads)} '# ' sections found; the outline in plan.json has more.")
    for b in cfg["blocks"]:
        if b["type"] == "image" and not (story_dir / b["file"]).exists():
            problems.append(f"Figure file '{b['file']}' does not exist. Use only: "
                            + ", ".join(f["file"] for f in plan.get("figures", [])))
        if b["type"] == "table":
            if not b["rows"]:
                problems.append(f"Table with columns {b['columns']} has no rows.")
            if len(b["columns"]) > 6:
                problems.append(f"Table with columns {b['columns']} has more than 6 columns; it will not fit the page.")
    images = {b["file"] for b in cfg["blocks"] if b["type"] == "image"}
    for f in plan.get("figures", []):
        if f["file"] not in images and (story_dir / f["file"]).exists():
            problems.append(f"Figure '{f['file']}' from plan.json is not placed in the document.")
    lowered = [h.lower() for h in heads]
    for a in plan.get("asks", []):
        sec = a.get("section", "")
        num = sec.split(" ", 1)[0].rstrip(".")
        if sec and num.isdigit() and not any(h.startswith(num + ".") for h in lowered):
            problems.append(f"Planned section '{sec}' (for ask: {a.get('ask', '')[:60]}) is missing.")
    if re.search(r"<(document title|title|subtitle|footer|caption)>", text):
        problems.append("Template placeholders like <title> are still in the text.")
    for h in heads:
        if re.fullmatch(r"\d+[.)]?\s*(figures?|visuals?|diagrams?)", h.strip(), re.I):
            problems.append(f"Section '{h}' only collects figures. Move each figure into the section it explains and delete this section.")
    if not problems:
        (story_dir / "content.json").write_text(json.dumps(cfg, indent=2, ensure_ascii=False), encoding="utf-8")
    return problems


def load_plan(story_dir):
    p = story_dir / "plan.json"
    if not p.exists():
        raise StepError("plan.json is missing – run the plan step first.")
    try:
        plan = json.loads(p.read_text(encoding="utf-8"))
    except json.JSONDecodeError as e:
        raise StepError(f"plan.json is not valid JSON: {e}")
    return plan


def plan_problems(plan, facts=None):
    req = ["story_id", "doc_title", "subtitle", "footer", "output_name", "asks", "outline", "sections", "figures", "evidence_queries"]
    probs = [f"missing key '{k}'" for k in req if k not in plan]
    if not probs:
        if not re.fullmatch(r"[A-Za-z0-9_\-]+", plan["output_name"]):
            probs.append("output_name must use letters, digits, '_' or '-' only")
        if not plan["asks"]:
            probs.append("asks is empty")
        for f in plan["figures"]:
            if not re.fullmatch(r"[A-Za-z0-9_\-]+\.png", f.get("file", "")):
                probs.append(f"figure file name '{f.get('file')}' must be a simple .png name")
        for q in plan["evidence_queries"]:
            if not q.get("question") or not q.get("keywords"):
                probs.append(f"evidence query {q.get('id')} needs a question and keywords")
        probs += plan_quality_problems(plan, facts or [])
    return probs


def plan_quality_problems(plan, facts):
    """Checks that make the plan deep enough: section types, fact assignments, no merged or figures-only sections."""
    from quality import patterns as K2
    probs = []
    types = set(K2.load_patterns()["types"]) | {"other"}
    heads = [s.get("heading", "") for s in plan.get("sections", [])]
    if heads != plan.get("outline", []):
        probs.append("'sections' must list the outline headings exactly, in the same order")
    for s in plan.get("sections", []):
        if s.get("type") not in types:
            probs.append(f"section '{s.get('heading')}' has type '{s.get('type')}'; use one of: {', '.join(sorted(types))}")
        if len(K2.classify(s.get("heading", ""))) > 1:
            probs.append(f"section '{s.get('heading')}' merges several topics ({', '.join(K2.classify(s.get('heading', '')))}); "
                         "give each its own section")
        if re.fullmatch(r"\d+[.)]?\s*(figures?|visuals?|diagrams?)", s.get("heading", "").strip(), re.I):
            probs.append(f"'{s.get('heading')}' is a figures-only section; place figures in the sections they explain")
    if facts:
        ids = {f["id"] for f in facts}
        assigned = {fid for s in plan.get("sections", []) for fid in s.get("facts", [])}
        unused = {u.get("id") for u in plan.get("unused_facts", []) if u.get("reason")}
        missing = sorted(ids - assigned - unused)
        if missing:
            probs.append(f"{len(missing)} facts are neither assigned to a section nor listed in unused_facts with a reason: "
                         + ", ".join(missing[:25]))
        if len(unused) > 0.15 * len(ids):
            probs.append(f"{len(unused)} of {len(ids)} facts are marked unused; keep it under 15%")
        bad = sorted(assigned - ids)
        if bad:
            probs.append("unknown fact IDs in sections: " + ", ".join(bad[:15]))
    if not any("(decision)" in h.lower() for h in heads) and re.search(r"open question|agree", plan.get("_story", ""), re.I):
        probs.append("the story has an open question or asks for agreement, but no section heading is marked (DECISION)")
    return probs


def write_section_types(story_dir):
    from quality import patterns as K2
    pats = K2.load_patterns()["types"]
    lines = ["# Section types", "", "Each section in plan.json gets one of these types. The draft is checked against the "
             "type's checklist and requirement floor.", ""]
    for t in pats:
        lines += [f"## {t}", "", K2.brief(t, n_exemplars=0), ""]
    (story_dir / "section_types.md").write_text("\n".join(lines), encoding="utf-8")


def write_section_briefs(story_dir, plan, facts):
    """section_briefs.md: for each planned section, its checklist, floor, assigned facts and gold exemplar rows."""
    from quality import patterns as K2
    by_id = {f["id"]: f for f in facts}
    figs = {}
    for f in plan.get("figures", []):
        figs.setdefault(f.get("section", ""), []).append(f)
    out = ["# Section briefs", "", f"Write each section of {plan.get('doc_title', 'the document')} to its brief. "
           "Use every assigned fact with its exact values. Meet the requirement floor. Cover every topic.", ""]
    for s in plan.get("sections", []):
        out += [f"## {s['heading']}", ""]
        if s.get("notes"):
            out += [f"Purpose: {s['notes']}", ""]
        t = s.get("type")
        if t and t != "other":
            out += [K2.brief(t), ""]
        if s.get("facts"):
            out.append("Assigned facts:")
            out += [f"- {fid}: {by_id[fid]['text']}" for fid in s["facts"] if fid in by_id]
            out.append("")
        for f in figs.get(s["heading"], []):
            out += [f"Figure here: {f['file']} – {f.get('caption', '')}", ""]
    path = story_dir / "section_briefs.md"
    path.write_text("\n".join(out) + "\n", encoding="utf-8")
    return path


# ------------------------------------------------------------------ steps
def step_check(run, args):
    problems = []
    for mod, pkg in (("docx", "python-docx"), ("matplotlib", "matplotlib"), ("pypdf", "pypdf")):
        try:
            __import__(mod)
        except ImportError:
            problems.append(f"Python package '{pkg}' is missing: python -m pip install -r requirements.txt")
    exe = os.environ.get("CODEX_BIN") or shutil.which("codex")
    if not exe and not run.dry:
        problems.append("Codex CLI not found on PATH (install it and sign in once, or set CODEX_BIN).")
    else:
        print(f"    codex: {exe or '(dry run – not needed)'}")
    if problems:
        raise StepError("; ".join(problems))
    return [], {}, "ready"


def step_setup(run, args):
    d = run.dir
    files = []
    if args.story:
        shutil.copy2(args.story, d / "story.md")
    if not (d / "story.md").exists():
        raise StepError("No story: pass --story <file with the user story prompt>.")
    files.append(d / "story.md")
    if args.research:
        shutil.copy2(args.research, d / "research_findings.md")
    if not (d / "research_findings.md").exists():
        raise StepError("No research findings: pass --research <research_findings.md>.")
    files.append(d / "research_findings.md")
    note = []
    src_dir = d / "research" / "sources"
    if args.sources:
        n = copy_tree(Path(args.sources), src_dir)
        note.append(f"{n} source file(s) copied")
    papers = [f for f in src_dir.glob("*") if f.suffix.lower() in (".pdf", ".html", ".htm", ".txt")] if src_dir.exists() else []
    note.append(f"{len(papers)} source paper(s) available")
    if args.inputs:
        copy_tree(Path(args.inputs), d / "inputs")
    converted = 0
    if (d / "inputs").exists():
        for f in sorted((d / "inputs").glob("*")):
            if f.is_file() and to_text(f, d / "inputs" / "_text"):
                converted += 1
    note.append(f"{converted} supporting input(s) converted to text")
    if args.example and Path(args.example).exists():
        shutil.copy2(args.example, d / "example_requirements.md")
        files.append(d / "example_requirements.md")
        note.append("example document copied")
    elif not (d / "example_requirements.md").exists():
        (d / "example_requirements.md").write_text("(No example document provided.)\n", encoding="utf-8")
        note.append("no example document")
    gold = Path(args.gold) if args.gold else None
    if gold and gold.exists():
        shutil.copy2(gold, d / "gold.md")
        note.append(f"gold document {gold.name}")
    write_section_types(d)
    files.append(d / "section_types.md")
    return files, {}, "; ".join(note)


def step_facts(run, args):
    from quality import facts as K1
    d = run.dir
    extra = [(p.name, p.read_text(encoding="utf-8", errors="ignore")) for p in sorted((d / "inputs" / "_text").glob("*.txt"))] \
        if (d / "inputs" / "_text").exists() else []
    facts, sources = K1.build((d / "research_findings.md").read_text(encoding="utf-8"), extra,
                              use_slm=args.critic_mode == "hybrid")
    j, m = K1.write(facts, sources, d)
    return [j, m, d / "sources.json"], {}, f"{len(facts)} facts, {len(sources)} key sources"


def step_plan(run, args):
    from quality import facts as K1
    tokens = {}
    if not args.validate_only:
        msg, t = run.codex("plan", run.prompt("1_plan.md", EXAMPLE=rel(run.dir / "example_requirements.md")))
        add_tokens(tokens, t)
    if run.dry:
        return [], tokens, "prompt written"
    facts, _ = K1.load(run.dir)
    story = (run.dir / "story.md").read_text(encoding="utf-8")
    for attempt in range(MAX_FIX_ROUNDS + 1):
        try:
            plan = load_plan(run.dir)
            probs = plan_problems(dict(plan, _story=story), facts)
        except StepError as e:
            probs = [str(e)]
        if not probs:
            briefs = write_section_briefs(run.dir, plan, facts)
            n_assigned = len({f for s in plan["sections"] for f in s.get("facts", [])})
            return [run.dir / "plan.json", briefs], tokens, (f"{len(plan['outline'])} sections, {len(plan['figures'])} figure(s), "
                                                             f"{len(plan['asks'])} asks, {n_assigned}/{len(facts)} facts assigned")
        if attempt == MAX_FIX_ROUNDS or args.validate_only:
            break
        print("    plan.json problems: " + "; ".join(probs))
        _, t = run.codex(f"plan_fix{attempt + 1}", f"PIPELINE STEP: plan (fix)\n\n{rel(run.dir)}/plan.json has these problems:\n- "
                         + "\n- ".join(probs) + "\n\nFix the file so it follows the structure and rules in prompts/1_plan.md "
                         f"(facts are in {rel(run.dir)}/facts.md). Change no other file.")
        add_tokens(tokens, t)
    raise StepError("plan.json still invalid: " + "; ".join(probs))


def step_evidence(run, args):
    plan = load_plan(run.dir) if not run.dry else {"evidence_queries": []}
    research = run.dir / "research"
    research.mkdir(exist_ok=True)
    src = research / "sources"
    if not src.exists() or not any(f.suffix.lower() in (".pdf", ".html", ".htm", ".txt") for f in src.glob("*")):
        (research / "research_evidence.md").write_text(
            "# Research Evidence\n\nNo source papers were provided (research/sources is empty). "
            "Use research_findings.md only, and mark facts that need a source as 'to be confirmed'.\n", encoding="utf-8")
        return [research / "research_evidence.md"], {}, "no source papers – evidence file says so"
    if run.dry:
        return [], {}, "skipped in dry run"
    questions = {"title": f"Research Evidence – {plan.get('doc_title', '')}",
                 "purpose": f"Evidence for {plan.get('subtitle', '')}",
                 "questions": [{"id": q.get("id", f"Q{i + 1}"), "heading": f"{i + 1}. {q.get('heading', q.get('id'))}",
                                "question": q["question"], "keywords": q["keywords"]}
                               for i, q in enumerate(plan["evidence_queries"])]}
    (research / "questions.json").write_text(json.dumps(questions, indent=2, ensure_ascii=False), encoding="utf-8")
    proc = subprocess.run([sys.executable, str(ROOT / "research_findings.py"), "--workdir", str(run.dir), "--top", "8"],
                          cwd=ROOT, capture_output=True, text=True)
    (run.dir / "logs" / "evidence.log").write_text(proc.stdout + proc.stderr, encoding="utf-8")
    if proc.returncode != 0:
        raise StepError("research_findings.py failed: " + (proc.stderr or proc.stdout).strip()[-400:])
    n = len(re.findall(r"^- \[", (research / "research_evidence.md").read_text(encoding="utf-8"), re.M))
    return [research / "questions.json", research / "research_evidence.md"], {}, f"{n} cited passages for {len(questions['questions'])} topics"


def step_figures(run, args):
    tokens = {}
    _, t = run.codex("figures", run.prompt("2_figures.md"))
    add_tokens(tokens, t)
    if run.dry:
        return [], tokens, "prompt written"
    plan = load_plan(run.dir)
    script = run.dir / "figures.py"
    qa_json = run.dir / "figure_qa.json"
    defects_left = ""
    for attempt in range(MAX_FIX_ROUNDS + 1):
        err, defects = "", ""
        if not script.exists():
            err = "figures.py was not created."
        else:
            # run the script under the C5 probe: it draws the PNGs and measures every label and arrow
            proc = subprocess.run([sys.executable, str(ROOT / "quality" / "figure_qa.py"), str(script), "--out", str(qa_json)]
                                  + (["--vlm", "--plan", str(run.dir / "plan.json")] if os.environ.get("SLM_VLM_MODEL") else []),
                                  cwd=run.dir, capture_output=True, text=True)
            (run.dir / "logs" / f"figure_qa_{attempt}.log").write_text(proc.stdout + proc.stderr, encoding="utf-8")
            if proc.returncode == 1:
                err = "figures.py failed:\n" + (proc.stderr or proc.stdout).strip()[-1500:]
            elif proc.returncode == 2:
                rep = json.loads(qa_json.read_text(encoding="utf-8"))
                defects = "\n".join(f"- {name}: [{d['severity']}] {d['defect']}: {d['detail']}"
                                    for name, v in rep["figures"].items() for d in v["defects"] if d["severity"] == "high")
        missing = [f["file"] for f in plan["figures"] if not (run.dir / f["file"]).exists()]
        if not err and missing:
            err = "These figure files were not created: " + ", ".join(missing)
        if not err and not defects:
            return [script, qa_json] + [run.dir / f["file"] for f in plan["figures"]], tokens, \
                f"{len(plan['figures'])} figure(s) drawn, no layout defects (after {attempt} fix round(s))"
        if attempt == MAX_FIX_ROUNDS:
            if err:
                break
            defects_left = defects
            return [script, qa_json] + [run.dir / f["file"] for f in plan["figures"]], tokens, \
                f"{len(plan['figures'])} figure(s) drawn; {defects_left.count(chr(10)) + 1} layout defect(s) remain (see figure_qa.json)"
        print(f"    figure problem, asking Codex to fix (round {attempt + 1})")
        body = err or ("The figures were drawn, but the layout check measured these defects:\n" + defects +
                       "\n\nFix each one: wrap long labels with \\n or widen the box, move overlapping text apart, "
                       "and route arrows in the gaps between boxes (around them, not through labels).")
        _, t = run.codex(f"figures_fix{attempt + 1}", f"PIPELINE STEP: figures (fix)\n\n{body}\n\nFix {rel(script)} so it runs "
                         "and saves every figure listed in plan.json. Follow prompts/2_figures.md. Change no other file.")
        add_tokens(tokens, t)
    raise StepError(err)


def draft_by_sections(run, args, tokens):
    """One Codex call per section (each with its own brief), then assemble content.md."""
    plan = load_plan(run.dir)
    briefs = (run.dir / "section_briefs.md").read_text(encoding="utf-8")
    blocks = {m.group(1).strip(): m.group(2) for m in re.finditer(r"(?ms)^## (.+?)\n(.*?)(?=^## |\Z)", briefs)}
    sec_dir = run.dir / "sections"
    sec_dir.mkdir(exist_ok=True)
    files = []
    for i, s in enumerate(plan["sections"], 1):
        name = f"{i:02d}_{re.sub(r'[^A-Za-z0-9]+', '_', s['heading']).strip('_')[:40]}.md"
        files.append(sec_dir / name)
        if (sec_dir / name).exists() and not args.redo_sections:
            continue
        _, t = run.codex(f"draft_s{i:02d}", run.prompt("8_section_draft.md", SECTION_NO=i, SECTION_TOTAL=len(plan["sections"]),
                                                       SECTION_FILE=name, SECTION_HEADING=s["heading"],
                                                       SECTION_BRIEF=blocks.get(s["heading"], "(no brief)"),
                                                       EXAMPLE=rel(run.dir / "example_requirements.md")))
        add_tokens(tokens, t)
    if run.dry:
        return
    parts = ["---", f"title: {plan['doc_title']}", f"subtitle: {plan['subtitle']}", f"footer: {plan['footer']}", "---", ""]
    for f in files:
        if not f.exists():
            raise StepError(f"section file {rel(f)} was not written")
        parts += [f.read_text(encoding="utf-8").strip(), ""]
    (run.dir / "content.md").write_text("\n".join(parts) + "\n", encoding="utf-8")


def step_draft(run, args):
    tokens = {}
    if args.draft_mode == "sections":
        draft_by_sections(run, args, tokens)
    else:
        _, t = run.codex("draft", run.prompt("3_draft.md", EXAMPLE=rel(run.dir / "example_requirements.md")))
        add_tokens(tokens, t)
    if run.dry:
        (run.dir / "_prompts" / "draft_fix_template.md").write_text(
            run.prompt("4_fix.md", ERRORS="(the pipeline fills in the problems it finds)"), encoding="utf-8")
        return [], tokens, "prompt written"
    plan = load_plan(run.dir)
    for attempt in range(MAX_FIX_ROUNDS + 1):
        problems = check_content(run.dir, plan)
        if not problems:
            return [run.dir / "content.md", run.dir / "content.json"], tokens, f"passed checks after {attempt} fix round(s)"
        if attempt == MAX_FIX_ROUNDS:
            break
        print(f"    content.md has {len(problems)} problem(s), asking Codex to fix (round {attempt + 1})")
        _, t = run.codex(f"draft_fix{attempt + 1}", run.prompt("4_fix.md", ERRORS="\n".join(f"- {p}" for p in problems)))
        add_tokens(tokens, t)
    raise StepError("content.md still has problems: " + "; ".join(problems))


def _findings_for_prompt(findings, limit=45):
    groups = {}
    for f in findings:
        if f["severity"] == "low":
            continue
        groups.setdefault(f["section"] or "(whole document)", []).append(f)
    lines, n = [], 0
    for sec, fs in sorted(groups.items(), key=lambda kv: min({"high": 0, "medium": 1}.get(f["severity"], 2) for f in kv[1])):
        lines.append(f"### {sec}")
        for f in fs:
            if n >= limit:
                break
            facts = f" (facts: {', '.join(f['suggested_facts'])})" if f.get("suggested_facts") else ""
            lines.append(f"- [{f['severity']}] {f['defect']} – {f['target']} – {f['evidence'][:220]} → {f['suggestion']}{facts}")
            n += 1
    rest = sum(1 for f in findings if f["severity"] != "low") - n
    if rest > 0:
        lines.append(f"\n({rest} more medium findings are in findings.jsonl.)")
    return "\n".join(lines)


def score_draft(run, args, label):
    """Run the critic panel + gate on content.md; write scorecard.json, findings.jsonl, critique_report.md."""
    from quality import critics as C
    from quality import scorecard as S
    d = run.dir
    gold = d / "gold.md" if (d / "gold.md").exists() else None
    ctx = C.context_for_story(d, gold_path=gold, mode=args.critic_mode)
    fig = json.loads((d / "figure_qa.json").read_text(encoding="utf-8")) if (d / "figure_qa.json").exists() else None
    card, findings = S.evaluate(ctx, fig_report=fig, judge=args.judge)
    card["label"] = label
    (d / "scorecard.json").write_text(json.dumps(card, indent=2, ensure_ascii=False), encoding="utf-8")
    with (d / "findings.jsonl").open("w", encoding="utf-8") as f:
        for x in findings:
            f.write(json.dumps(x, ensure_ascii=False) + "\n")
    md = S.to_markdown(card, findings, f"Quality scorecard – {label}")
    (d / "critique_report.md").write_text(md, encoding="utf-8")
    (d / "logs" / f"critique_{re.sub(r'[^A-Za-z0-9]+', '_', label)}.md").write_text(md, encoding="utf-8")
    return card, findings


def step_critique(run, args):
    """Critic panel → gate → targeted revision (R1), up to --revise-rounds times."""
    tokens = {}
    if run.dry:
        (run.dir / "_prompts" / "revise_template.md").write_text(
            run.prompt("6_revise.md", ROUND=1, GATE_SUMMARY="(filled in by the pipeline)",
                       FINDINGS="(the critic findings, grouped by section)"), encoding="utf-8")
        return [], tokens, "revise prompt template written"
    plan = load_plan(run.dir)
    history = []
    card, findings = score_draft(run, args, "round 0")
    history.append(card)
    for rnd in range(1, 0 if args.validate_only else args.revise_rounds + 1):
        if card["pass"]:
            break
        summary = ("hard gates failed: " + ", ".join(card["hard_failures"]) if card["hard_failures"] else "hard gates pass") + \
            (f"; score vs gold {card['relative_to_gold']:.0%}" if card.get("relative_to_gold") is not None else "")
        print(f"    gate failed ({summary}); revising (round {rnd})")
        _, t = run.codex(f"revise{rnd}", run.prompt("6_revise.md", ROUND=rnd, GATE_SUMMARY=summary,
                                                    FINDINGS=_findings_for_prompt(findings)))
        add_tokens(tokens, t)
        problems = check_content(run.dir, plan)
        if problems:
            _, t = run.codex(f"revise{rnd}_fix", run.prompt("4_fix.md", ERRORS="\n".join(f"- {p}" for p in problems)))
            add_tokens(tokens, t)
            problems = check_content(run.dir, plan)
            if problems:
                raise StepError("content.md has format problems after revision: " + "; ".join(problems))
        card, findings = score_draft(run, args, f"round {rnd}")
        history.append(card)
    if not args.validate_only:
        shutil.copy2(run.dir / "content.md", run.dir / "content.generated.md")   # snapshot for the feedback flywheel
    trail = " → ".join(f"{(h['relative_to_gold'] or h['overall']):.0%}" for h in history)
    note = (f"gate {'PASS' if card['pass'] else 'FAIL'} after {len(history) - 1} revision round(s); score {trail}; "
            f"{card['counts']['high']} high findings left")
    if not card["pass"] and args.strict_gates:
        raise StepError(note + " – see critique_report.md")
    return [run.dir / "scorecard.json", run.dir / "findings.jsonl", run.dir / "critique_report.md"], tokens, note


def step_build(run, args):
    if run.dry:
        return [], {}, "skipped in dry run"
    import generate_doc
    plan = load_plan(run.dir)
    cfg = json.loads((run.dir / "content.json").read_text(encoding="utf-8"))
    name = plan["output_name"]
    generate_doc.render_docx(cfg, run.dir, run.dir / f"{name}.docx")
    generate_doc.render_md(cfg, run.dir / f"{name}.md")
    return [run.dir / f"{name}.docx", run.dir / f"{name}.md"], {}, "Word and Markdown built"


def step_qa(run, args):
    msg, t = run.codex("qa", run.prompt("5_qa.md", REVIEWER=args.reviewer), sandbox="read-only", out_name="qa_last_message.md")
    if run.dry:
        return [], t, "prompt written"
    out = run.dir / "qa_report.md"
    out.write_text(f"# QA Report – {load_plan(run.dir).get('doc_title', '')}\n\n_Generated {now()} by Codex (read-only review)._\n\n"
                   + (msg.strip() or "_Codex returned no review text – see logs/qa.log._") + "\n", encoding="utf-8")
    high_rows = [l for l in msg.splitlines() if re.search(r"\|\s*High\s*\|", l, re.I)]
    high = len(high_rows)
    if not high or args.no_qa_fix:
        return [out], t, f"{high} high-severity issue(s)"
    # blocking QA: one fix round for the High issues, then re-check, re-score and rebuild
    print(f"    {high} High QA issue(s); asking Codex to fix them")
    _, t2 = run.codex("qa_fix", run.prompt("4_fix.md", ERRORS="These High issues were found by the QA review "
                                           "(see qa_report.md):\n" + "\n".join(high_rows)))
    add_tokens(t, t2)
    plan = load_plan(run.dir)
    problems = check_content(run.dir, plan)
    if problems:
        raise StepError("content.md has format problems after the QA fix: " + "; ".join(problems))
    card, _ = score_draft(run, args, "after QA fix")
    step_build(run, args)
    with out.open("a", encoding="utf-8") as f:
        f.write(f"\n## After the QA fix\n\nThe {high} High issue(s) above were sent back for one fix round; the document was "
                f"rebuilt. Quality gate after the fix: {'PASS' if card['pass'] else 'FAIL'}.\n")
    return [out], t, f"{high} high-severity issue(s) sent for one fix round; gate {'PASS' if card['pass'] else 'FAIL'} after it"


def step_report(run, args):
    if run.dry:
        return [], {}, "skipped in dry run"
    import generate_doc
    import md_to_content
    d = run.dir
    plan = load_plan(d)
    steps = run.state["steps"]
    total = {}
    for s in steps.values():
        add_tokens(total, s.get("tokens"))
    tok_in = total.get("input_tokens", 0)
    tok_out = total.get("output_tokens", 0)
    content = (d / "content.md").read_text(encoding="utf-8") if (d / "content.md").exists() else ""
    tbc = [re.sub(r"\s+", " ", re.sub(r"^[\s>#*|\-]+", "", l)).strip()[:160]
           for l in content.splitlines() if "to be confirmed" in l.lower()]
    cites = sorted(set(re.findall(r"\[([^\]\[]+?\.(?:pdf|html?|txt))(?:, p\.\d+)?\]", content)))
    papers = sorted(f.name for f in (d / "research" / "sources").glob("*")
                    if f.suffix.lower() in (".pdf", ".html", ".htm", ".txt")) if (d / "research" / "sources").exists() else []
    out_doc = d / f"{plan['output_name']}.docx"
    lines = ["---", f"title: Run Report – {plan['doc_title']}", f"subtitle: {plan['subtitle']} · generated {dt.date.today()}",
             f"footer: Run report · {plan['footer']}", "---", "",
             "# 1. Summary", "",
             f"The pipeline turned the user story in story.md and the research in research_findings.md into {out_doc.name}, "
             f"using {len(papers)} source paper(s). Codex wrote the plan, figures, document text and QA review; the project's "
             "Python scripts extracted evidence, checked the text and built the Word file. Review the QA issues and the "
             "items marked 'to be confirmed' before publishing.", "",
             "| Item | Value |", "|---|---|",
             f"| **Document** | {out_doc.name} |",
             f"| **Story folder** | {rel(d)} |",
             f"| **Sections / figures** | {len(plan['outline'])} / {len(plan['figures'])} |",
             f"| **Source papers / cited in text** | {len(papers)} / {len(cites)} |",
             f"| **Items marked to be confirmed** | {len(tbc)} |",
             f"| **Codex tokens (input / output)** | {tok_in:,} / {tok_out:,} |" if total else "| **Codex tokens** | not reported by Codex |",
             "", "# 2. Steps", "", "| Step | Status | Time (s) | Result |", "|---|---|---|---|"]
    for s in STEPS:
        if s in steps and s != "report":
            st = steps[s]
            lines.append(f"| **{s}** | {st['status']} | {st['seconds']} | {st.get('note', '').replace('|', '/')} |")
    lines += ["", "# 3. Files produced", "", "| File | What it is |", "|---|---|"]
    desc = {"plan.json": "Document plan written by Codex", "content.md": "Document text (Markdown) written by Codex",
            "content.json": "Document text converted for the Word builder", "figures.py": "Figure script written by Codex",
            "qa_report.md": "Codex fact-check and coverage review", "research_evidence.md": "Cited passages from the source papers",
            "questions.json": "Evidence topics from the plan", "story.md": "The user story (input)",
            "research_findings.md": "Research findings (input)",
            "example_requirements.md": "Example document used as the quality reference (input)",
            "run_report.md": "This report", "run_report.docx": "This report (Word)",
            "facts.jsonl": "Fact ledger (machine-readable)", "facts.md": "Fact ledger", "sources.json": "Key sources from the findings",
            "section_briefs.md": "Per-section briefs given to the writer", "section_types.md": "Section types and checklists",
            "scorecard.json": "Quality scorecard (D1–D10, gates)", "findings.jsonl": "Critic findings",
            "critique_report.md": "Quality scorecard and findings", "figure_qa.json": "Measured figure layout defects"}
    seen = set()
    for s in STEPS:
        for f in steps.get(s, {}).get("files", []):
            if f in seen:
                continue
            seen.add(f)
            name = Path(f).name
            what = desc.get(name, "Figure" if name.endswith(".png") else
                            "Requirements document" if name.startswith(plan["output_name"]) else "")
            lines.append(f"| {name} | {what} |")
    lines += ["", "# 4. Story coverage", "", "| Ask | Section |", "|---|---|"]
    lines += [f"| {a.get('ask', '').replace('|', '/')} | {a.get('section', '')} |" for a in plan["asks"]]
    lines += ["", "# 5. Items marked to be confirmed", ""]
    lines += [f"- {t}" for t in tbc[:40]] or ["- None."]
    lines += ["", "# 6. QA review", ""]
    qa = d / "qa_report.md"
    if qa.exists():
        body = qa.read_text(encoding="utf-8").split("\n", 1)[-1]
        body = re.sub(r"^#{1,6}\s*", "## ", body, flags=re.M)
        lines += [l for l in body.splitlines() if not l.startswith("_Generated")]
    else:
        lines += ["- QA step was not run."]
    lines += ["", "# 7. Quality scorecard", ""]
    sc = d / "scorecard.json"
    if sc.exists():
        card = json.loads(sc.read_text(encoding="utf-8"))
        lines += [f"Gate: {'PASS' if card['pass'] else 'FAIL'}"
                  + (f" – hard gates failed: {', '.join(card['hard_failures'])}" if card["hard_failures"] else "")
                  + (f". Score against the gold document: {card['relative_to_gold']:.0%}." if card.get('relative_to_gold') is not None else "."),
                  "", "| Dimension | Score | Gold | Gate |", "|---|---|---|---|"]
        gd = card.get("gold_dimensions") or {}
        for k, v in card["dimensions"].items():
            g = card["gates"].get(k)
            lines.append(f"| **{k}** | {'–' if v is None else v} | {'–' if gd.get(k) is None else gd[k]} | "
                         f"{'–' if not g else ('pass' if g['pass'] else 'FAIL' if g['hard'] else 'below target')} |")
        lines += ["", "> Findings are in critique_report.md and findings.jsonl. Edit content.md, then run --from build; your "
                  "edits are compared with content.generated.md to train the critics (python quality/training.py export)."]
    else:
        lines += ["- The critique step was not run."]
    lines += ["", "# 8. Next steps", "",
              "- Fix the high-severity QA issues in content.md, then rebuild: python run_pipeline.py --id "
              f"{plan['story_id']} --from build",
              "- Check every item marked to be confirmed with the owner or the source document.",
              "- Read the Word document in full, then upload it to the team file share."]
    report_md = d / "run_report.md"
    report_md.write_text("\n".join(lines) + "\n", encoding="utf-8")
    cfg = md_to_content.parse(report_md.read_text(encoding="utf-8"))
    generate_doc.render_docx(cfg, d, d / "run_report.docx")
    return [report_md, d / "run_report.docx"], {}, "run report written"


STEP_FUNCS = {"check": step_check, "setup": step_setup, "facts": step_facts, "plan": step_plan, "evidence": step_evidence,
              "figures": step_figures, "draft": step_draft, "critique": step_critique, "build": step_build,
              "qa": step_qa, "report": step_report}


# ------------------------------------------------------------------ main
def story_id_from(path):
    text = Path(path).read_text(encoding="utf-8", errors="ignore")
    m = re.search(r"#?\b(\d{5,8})\b", text)
    return m.group(1) if m else Path(path).stem


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--story", help="file with the user story prompt")
    ap.add_argument("--research", help="research_findings.md")
    ap.add_argument("--sources", help="folder with source papers (PDF/HTML/TXT)")
    ap.add_argument("--inputs", help="folder with other supporting files")
    ap.add_argument("--reviewer", default=os.environ.get("QA_REVIEWER", "subject-matter expert for the topic of the story"),
                    help="who the QA step should review as, e.g. \"data platform architect\"")
    ap.add_argument("--example", default=str(DEFAULT_GOLD),
                    help="a finished, reviewed requirements document (Markdown) used as the quality and style reference "
                         "(default: the US311617 gold document)")
    ap.add_argument("--gold", default=str(DEFAULT_GOLD), help="gold document the draft is scored against")
    ap.add_argument("--draft-mode", choices=["whole", "sections"], default=os.environ.get("DRAFT_MODE", "whole"),
                    help="whole = one Codex call; sections = one call per section with its own brief")
    ap.add_argument("--redo-sections", action="store_true", help="with --draft-mode sections: rewrite sections that already exist")
    ap.add_argument("--revise-rounds", type=int, default=int(os.environ.get("REVISE_ROUNDS", "2")),
                    help="critique → revise rounds when the quality gate fails (default 2)")
    ap.add_argument("--strict-gates", action="store_true", help="fail the run if hard gates still fail after revision")
    ap.add_argument("--critic-mode", choices=["rules", "hybrid"], default=os.environ.get("CRITIC_MODE", "rules"),
                    help="hybrid = rules + small model critics (needs SLM_BASE_URL / SLM_MODEL)")
    ap.add_argument("--judge", action="store_true", help="add the J1 pairwise judge (needs JUDGE_* or LLM_* settings)")
    ap.add_argument("--no-qa-fix", action="store_true", help="keep QA advisory (no automatic fix round for High issues)")
    ap.add_argument("--validate-only", action="store_true",
                    help="with --only plan: check plan.json and write section_briefs.md without calling Codex; "
                         "with --only critique: score content.md without revising")
    ap.add_argument("--id", help="story id (default: first 5–8 digit number in the story)")
    ap.add_argument("--from", dest="start", choices=STEPS, help="start at this step")
    ap.add_argument("--only", choices=STEPS, help="run only this step")
    ap.add_argument("--dry-run", action="store_true", help="write the Codex prompts to _prompts/ without calling Codex")
    ap.add_argument("--codex-args", default=os.environ.get("CODEX_ARGS", ""), help="extra arguments for codex exec")
    args = ap.parse_args()

    if not args.id and not args.story:
        ap.error("give --story (new run) or --id (existing story folder)")
    sid = args.id or story_id_from(args.story)
    story_dir = ROOT / "stories" / re.sub(r"[^A-Za-z0-9_\-]", "_", sid)
    story_dir.mkdir(parents=True, exist_ok=True)
    run = Run(story_dir, dry_run=args.dry_run, codex_args=shlex.split(args.codex_args))

    steps = [args.only] if args.only else STEPS[STEPS.index(args.start):] if args.start else STEPS
    if args.only and args.only != "check" and "setup" not in run.state["steps"] and args.only != "setup":
        steps = ["setup", args.only]
    print(f"Story {sid}  ->  {rel(story_dir)}{'   (DRY RUN: prompts only)' if args.dry_run else ''}")
    for s in steps:
        print(f"[{s}]")
        t0 = time.time()
        try:
            files, tokens, note = STEP_FUNCS[s](run, args)
        except StepError as e:
            run.record(s, "FAILED", time.time() - t0, note=str(e))
            print(f"    FAILED: {e}\n\nFix the problem, then continue with:  python run_pipeline.py --id {sid} --from {s}")
            if s != "report":
                try:
                    step_report(run, args)
                except Exception:
                    pass
            sys.exit(1)
        run.record(s, "ok" if not args.dry_run else "dry-run", time.time() - t0, files, tokens, note)
        tok = f"  tokens in/out {tokens.get('input_tokens', 0):,}/{tokens.get('output_tokens', 0):,}" if tokens else ""
        print(f"    ok – {note}{tok}")
    if args.dry_run:
        print(f"\nDry run done. The prompts Codex would receive are in {rel(story_dir / '_prompts')}.")
    else:
        plan = load_plan(story_dir)
        print(f"\nDone. Document: {rel(story_dir / (plan['output_name'] + '.docx'))}   Report: {rel(story_dir / 'run_report.docx')}")


if __name__ == "__main__":
    main()
